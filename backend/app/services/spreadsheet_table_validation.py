"""Read-only, deliberately narrow verification of cited XLSX table evidence.

This is an evidence extractor, not an answer validator. Callers must separately
check that every claim (including time and footnote qualifiers) is supported.
"""
from hashlib import sha256
from pathlib import Path
import re


def _text(value):
    return '' if value is None else str(value).strip()


def _cells(line):
    if not line.startswith('|') or not line.endswith('|'):
        return None
    return [part.strip() for part in line[1:-1].split('|')]


def verify_spreadsheet_cell(session, citation, kb_ids, *, source_sha256,
                            sheet, title_row, header_end_row, row, label_col, value_col):
    """Verify an explicitly located cell and return its source context or None.

    source_sha256 must come from trusted ingestion lineage, not a hash computed
    from the current file at verification time. All coordinates are 1-based.
    Supported layout: one title cell, contiguous multi-row headers, unique row
    label, a numeric literal cell, and footnotes after the table. The citation
    must contain the title, each nonempty header cell in its column, and the
    exact row label/value at their respective positions. No fuzzy matching.

    Returned mapping includes source identity/hash, coordinates, title,
    complete header fragments, literal row/value, FY scope, referenced
    footnotes and source/notes. It is not a decision that an answer is correct.
    Unsupported or ambiguous formats return None.
    """
    try:
        from app.models import DocumentChunk, FileDocument
        import openpyxl
        if (not re.fullmatch(r'[0-9a-f]{64}', source_sha256 or '') or
                not all(type(x) is int and x > 0 for x in
                        (title_row, header_end_row, row, label_col, value_col)) or
                not title_row < header_end_row < row or label_col == value_col):
            return None
        chunk = session.get(DocumentChunk, citation.get('chunk_id'))
        doc = session.get(FileDocument, citation.get('fileId')) if chunk else None
        if (not chunk or not doc or chunk.file_id != doc.id or doc.kb_id not in kb_ids
                or citation.get('kb_id') != doc.kb_id or chunk.content != citation.get('content')):
            return None
        content = chunk.content
        if not isinstance(content, str) or not content.startswith(f'[工作表: {sheet}]\n'):
            return None
        path = Path(doc.local_path)
        if path.suffix.lower() != '.xlsx' or not path.is_file() or path.stat().st_size > 50_000_000:
            return None
        before = path.stat()
        data = path.read_bytes()
        if sha256(data).hexdigest() != source_sha256:
            return None
        # Parse the exact bytes we hashed, so a concurrent replacement cannot
        # cause the proof to describe a different workbook.
        from io import BytesIO
        workbook = openpyxl.load_workbook(BytesIO(data), read_only=False, data_only=False)
        try:
            if sheet not in workbook.sheetnames:
                return None
            ws = workbook[sheet]
            if ws.max_row > 100000 or ws.max_column > 1000:
                return None
            if any(r.min_row <= row and r.max_row > title_row for r in ws.merged_cells.ranges):
                return None
            title = _text(ws.cell(title_row, 1).value)
            if not title or any(ws.cell(title_row, c).value is not None for c in range(2, ws.max_column + 1)):
                return None
            matches = re.findall(r'(?i)\bfiscal\s+year\s+(\d{4})\b', title)
            if len(matches) != 1:
                return None
            label = _text(ws.cell(row, label_col).value)
            cell = ws.cell(row, value_col)
            if (not label or cell.data_type == 'f' or type(cell.value) not in (int, float)
                    or not re.fullmatch(r'-?\d+(?:\.\d+)?', str(cell.value))):
                return None
            if sum(_text(ws.cell(r, label_col).value) == label for r in range(1, ws.max_row + 1)) != 1:
                return None
            headers = []
            for r in range(title_row + 1, header_end_row + 1):
                for c in (label_col, value_col):
                    v = ws.cell(r, c)
                    if v.data_type == 'f':
                        return None
                    if _text(v.value):
                        headers.append((r, c, _text(v.value)))
            if (not _text(ws.cell(header_end_row, label_col).value)
                    or not _text(ws.cell(header_end_row, value_col).value)):
                return None
            lines = [_cells(line) for line in content.splitlines()]
            lines = [line for line in lines if line]
            if not lines or lines[0][0] != title:
                return None
            # The chunk must represent the complete physical header rows in order.
            header_lines = []
            for r in range(title_row + 1, header_end_row + 1):
                expected = [_text(ws.cell(r, c).value) or 'nan' for c in range(1, len(lines[0]) + 1)]
                header_lines.append(expected)
            if lines[2:2 + len(header_lines)] != header_lines:
                return None
            if any(_text(ws.cell(r, label_col).value) or _text(ws.cell(r, value_col).value)
                   for r in range(header_end_row + 1, row) if _text(ws.cell(r, label_col).value) == label):
                return None
            row_lines = [line for line in lines[2 + len(header_lines):]
                         if len(line) >= max(label_col, value_col) and line[label_col - 1] == label]
            if len(row_lines) != 1 or row_lines[0][value_col - 1] != str(cell.value):
                return None
            # A second title between the declared title and row invalidates the locator.
            if any(re.match(r'(?i)^table\s+\d+\b', _text(ws.cell(r, 1).value))
                   for r in range(title_row + 1, row)):
                return None
            markers = set()
            for _, c, h in headers:
                match = re.search(r'(?<=[A-Za-z)])(\d+)$', h)
                if match:
                    markers.add(match.group(1))
            notes = []
            source = []
            footnotes = {}
            for r in range(row + 1, ws.max_row + 1):
                t = _text(ws.cell(r, 1).value)
                if re.match(r'(?i)^table\s+\d+\b', t):
                    return None
                m = re.match(r'^(\d+)(\D.+)$', t, re.S)
                if m and m.group(1) in markers:
                    if m.group(1) in footnotes:
                        return None
                    footnotes[m.group(1)] = {'text': t, 'row': r, 'column': 1}
                if re.match(r'(?i)^note\s*:', t): notes.append(t)
                if re.match(r'(?i)^source\s*:', t): source.append(t)
            if set(footnotes) != markers or len(source) != 1:
                return None
            source_years = re.findall(r'(?i)\b(?:fy|fiscal\s+year)\s*(\d{4})\b', source[0])
            if not source_years or set(source_years) != set(matches):
                return None
            after = path.stat()
            if (before.st_size, before.st_mtime_ns, before.st_ino) != (after.st_size, after.st_mtime_ns, after.st_ino):
                return None
            return {'chunk_id': citation['chunk_id'], 'file_id': doc.id, 'kb_id': doc.kb_id,
                    'source_sha256': source_sha256, 'sheet': sheet,
                    'row': row, 'label_col': label_col, 'value_col': value_col,
                    'title': title, 'fiscal_year': matches[0], 'headers': headers,
                    'label': label, 'value': cell.value, 'footnotes': footnotes,
                    'notes': notes, 'source': source[0]}
        finally:
            workbook.close()
    except Exception:
        return None



def discover_spreadsheet_cells(session, citation, kb_ids, *, source_sha256):
    """Discover verified numeric cells in a single cited Markdown table.

    No query, expected answer or coordinate is accepted. The digest must come
    from trusted ingestion lineage. A successful result is a list of the same
    evidence mappings as verify_spreadsheet_cell; [] means unavailable or
    unsupported. The implementation intentionally caps workbook dimensions,
    citation size and candidates and refuses ambiguous title/header layouts.
    Discovery alone does not validate any answer or its qualifiers.
    """
    try:
        from app.models import DocumentChunk, FileDocument
        from io import BytesIO
        import openpyxl
        if not re.fullmatch(r'[0-9a-f]{64}', source_sha256 or ''):
            return []
        chunk = session.get(DocumentChunk, citation.get('chunk_id'))
        doc = session.get(FileDocument, citation.get('fileId')) if chunk else None
        content = citation.get('content')
        if (not chunk or not doc or chunk.file_id != doc.id or doc.kb_id not in kb_ids
                or citation.get('kb_id') != doc.kb_id or chunk.content != content
                or not isinstance(content, str) or len(content) > 250000):
            return []
        match = re.match(r'^\[工作表: ([^\]\r\n]+)\]\n', content)
        if not match:
            return []
        sheet = match.group(1)
        lines = [_cells(line) for line in content.splitlines()]
        lines = [line for line in lines if line]
        if (len(lines) < 4 or not 2 <= len(lines[0]) <= 32
                or any(len(line) != len(lines[0]) for line in lines)):
            return []
        path = Path(doc.local_path)
        if path.suffix.lower() != '.xlsx' or not path.is_file() or path.stat().st_size > 50_000_000:
            return []
        data = path.read_bytes()
        if sha256(data).hexdigest() != source_sha256:
            return []
        wb = openpyxl.load_workbook(BytesIO(data), read_only=False, data_only=False)
        try:
            if sheet not in wb.sheetnames:
                return []
            ws = wb[sheet]
            if ws.max_row > 2000 or ws.max_column > 128:
                return []
            titles = [r for r in range(1, ws.max_row + 1)
                      if _text(ws.cell(r, 1).value) == lines[0][0]]
            if len(titles) != 1:
                return []
            title_row = titles[0]
            width = len(lines[0])
            boundaries = []
            for end in range(title_row + 1, min(title_row + 30, ws.max_row) + 1):
                if not isinstance(ws.cell(end, 1).value, str):
                    continue
                cols = [col for col in range(2, width + 1)
                        if isinstance(ws.cell(end, col).value, str)
                        and _text(ws.cell(end, col).value)]
                expected = [[_text(ws.cell(r, col).value) or 'nan' for col in range(1, width + 1)]
                            for r in range(title_row + 1, end + 1)]
                if cols and lines[2:2 + len(expected)] == expected:
                    boundaries.append((end, cols))
            if len(boundaries) != 1:
                return []
            end, columns = boundaries[0]
            body = lines[2 + end - title_row:]
            labels = {}
            for r in range(end + 1, ws.max_row + 1):
                label = _text(ws.cell(r, 1).value)
                if label:
                    labels.setdefault(label, []).append(r)
            candidates = []
            for line in body:
                rows = labels.get(line[0], [])
                if len(rows) != 1:
                    continue
                row = rows[0]
                for col in columns:
                    cell = ws.cell(row, col)
                    if (cell.data_type != 'f' and type(cell.value) in (int, float)
                            and line[col - 1] == str(cell.value)):
                        candidates.append((row, col))
            if len(candidates) > 128 or len(set(candidates)) != len(candidates):
                return []
        finally:
            wb.close()
        results = []
        for row, col in candidates:
            proof = verify_spreadsheet_cell(session, citation, kb_ids, source_sha256=source_sha256,
                                             sheet=sheet, title_row=title_row, header_end_row=end,
                                             row=row, label_col=1, value_col=col)
            if proof is not None:
                results.append(proof)
        return results
    except Exception:
        return []
