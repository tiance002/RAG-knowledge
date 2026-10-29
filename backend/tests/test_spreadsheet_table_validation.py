import hashlib
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl
import pytest
from app.models import DocumentChunk, FileDocument
from app.services.spreadsheet_table_validation import verify_spreadsheet_cell, discover_spreadsheet_cells


def setup(path, content, digest=None):
    chunk = SimpleNamespace(file_id='file', content=content)
    doc = SimpleNamespace(id='file', kb_id='kb', local_path=str(path))
    session = MagicMock()
    session.get.side_effect = lambda model, ident: (chunk if model is DocumentChunk and ident == 'chunk'
        else doc if model is FileDocument and ident == 'file' else None)
    citation = {'chunk_id': 'chunk', 'fileId': 'file', 'kb_id': 'kb', 'content': content}
    options = dict(source_sha256=digest or hashlib.sha256(path.read_bytes()).hexdigest(), sheet='T01',
                   title_row=1, header_end_row=7, row=11, label_col=1, value_col=2)
    return session, citation, options


def test_saved_real_evidence_and_rejections():
    root = Path(__file__).resolve().parents[2]
    path = root / 'var/p0-observability/p1-live-review/excel-1.xlsx'
    response = root / 'var/p0-observability/p1-live-review/excel-response.sse'
    if not path.exists() or not response.exists():
        pytest.skip('local saved workbook and response are unavailable')
    line = next(x for x in response.read_text(encoding='utf8').splitlines() if x.startswith('data: {"citations"'))
    citation = json.loads(line[6:])['citations'][0]
    session, c, opts = setup(path, citation['content'], '6669cace46f876b0b09b67d91296c6a56e3268448c02f1a9e9dbceb4a80a43d9')
    proof = verify_spreadsheet_cell(session, c, ['kb'], **opts)
    assert proof is not None
    assert (proof['label'], proof['value'], proof['fiscal_year']) == ('Alabama', 219, '2018')
    assert [h[2] for h in proof['headers'] if h[1] == 2] == ['Number', 'of public', 'libraries1']
    assert 'administrative entity' in proof['footnotes']['1']['text']
    assert 'FY 2018' in proof['source']
    assert verify_spreadsheet_cell(session, {**c, 'content': c['content'] + 'x'}, ['kb'], **opts) is None
    assert verify_spreadsheet_cell(session, {**c, 'kb_id': 'other'}, ['kb'], **opts) is None
    assert verify_spreadsheet_cell(session, c, ['other'], **opts) is None
    assert verify_spreadsheet_cell(session, c, ['kb'], **{**opts, 'source_sha256': '0' * 64}) is None
    assert verify_spreadsheet_cell(session, c, ['kb'], **{**opts, 'value_col': 3})['value'] != 219
    assert verify_spreadsheet_cell(session, c, ['kb'], **{**opts, 'sheet': 'T02'}) is None


def test_synthetic_ambiguous_formula_footnote_and_stale(tmp_path):
    path = tmp_path / 'source.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'T01'
    ws['A1'] = 'Table 1. Counts: Fiscal year 2020'
    ws['B5'], ws['B6'], ws['A7'], ws['B7'] = 'Number', 'of public', 'State', 'libraries1'
    ws['A11'], ws['B11'] = 'Example', 12
    ws['A12'] = '1Administrative entities, not physical outlets.'
    ws['A13'] = 'SOURCE: Survey, FY 2020.'
    wb.save(path)
    content = ('[工作表: T01]\n| Table 1. Counts: Fiscal year 2020 | Unnamed: 1 |\n|---|---|\n'
               '| nan | nan |\n| nan | nan |\n| nan | nan |\n| nan | Number |\n'
               '| nan | of public |\n| State | libraries1 |\n| nan | nan |\n| Example | 12 |')
    session, c, opts = setup(path, content)
    assert verify_spreadsheet_cell(session, c, ['kb'], **opts)['footnotes']['1']['text'].startswith('1Administrative')
    ws['A14'], ws['B14'] = 'Example', 12
    wb.save(path)
    assert verify_spreadsheet_cell(session, c, ['kb'], **opts) is None  # stale hash
    opts['source_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    assert verify_spreadsheet_cell(session, c, ['kb'], **opts) is None  # duplicate entity
    ws['A14'], ws['B14'] = None, None
    ws['B11'] = '=6+6'
    wb.save(path)
    opts['source_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    assert verify_spreadsheet_cell(session, c, ['kb'], **opts) is None
    ws['B11'] = 12
    ws['A12'] = None
    wb.save(path)
    opts['source_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    assert verify_spreadsheet_cell(session, c, ['kb'], **opts) is None
    ws['A12'] = '1Administrative entities.'
    ws['A13'] = 'SOURCE: Survey, FY 2021.'
    wb.save(path)
    opts['source_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    assert verify_spreadsheet_cell(session, c, ['kb'], **opts) is None


def test_automatic_discovery_from_saved_citation():
    root = Path(__file__).resolve().parents[2]
    path = root / 'var/p0-observability/p1-live-review/excel-1.xlsx'
    response = root / 'var/p0-observability/p1-live-review/excel-response.sse'
    if not path.exists() or not response.exists():
        pytest.skip('local saved workbook and response are unavailable')
    line = next(x for x in response.read_text(encoding='utf8').splitlines() if x.startswith('data: {"citations"'))
    original = json.loads(line[6:])['citations'][0]
    session, citation, opts = setup(path, original['content'])
    proofs = discover_spreadsheet_cells(session, citation, ['kb'], source_sha256=opts['source_sha256'])
    assert any(p['label'] == 'Alabama' and p['row'] == 11 and p['value_col'] == 2
               and p['value'] == 219 and p['fiscal_year'] == '2018'
               and 'administrative entity' in p['footnotes']['1']['text'] for p in proofs)
    assert discover_spreadsheet_cells(session, citation, ['other'], source_sha256=opts['source_sha256']) == []
    assert discover_spreadsheet_cells(session, citation, ['kb'], source_sha256='0' * 64) == []


def test_automatic_discovery_synthetic_and_ambiguous(tmp_path):
    path = tmp_path / 'source.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'T01'
    ws['A1'] = 'Table 1. Counts: Fiscal year 2020'
    ws['B5'], ws['B6'], ws['A7'], ws['B7'] = 'Number', 'of public', 'State', 'libraries1'
    ws['A11'], ws['B11'] = 'Example', 12
    ws['A12'] = '1Administrative entities.'
    ws['A13'] = 'SOURCE: Survey, FY 2020.'
    wb.save(path)
    content = ('[工作表: T01]\n| Table 1. Counts: Fiscal year 2020 | Unnamed: 1 |\n|---|---|\n'
               '| nan | nan |\n| nan | nan |\n| nan | nan |\n| nan | Number |\n'
               '| nan | of public |\n| State | libraries1 |\n| nan | nan |\n| Example | 12 |')
    session, citation, opts = setup(path, content)
    result = discover_spreadsheet_cells(session, citation, ['kb'], source_sha256=opts['source_sha256'])
    assert len(result) == 1 and result[0]['row'] == 11 and result[0]['value_col'] == 2
    assert discover_spreadsheet_cells(session, {**citation, 'content': content + 'x'}, ['kb'], source_sha256=opts['source_sha256']) == []
    ws['A14'], ws['B14'] = 'Example', 12
    wb.save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    assert discover_spreadsheet_cells(session, citation, ['kb'], source_sha256=digest) == []
    ws['A14'], ws['B14'] = None, None
    ws['A15'] = 'Table 1. Counts: Fiscal year 2020'
    wb.save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    assert discover_spreadsheet_cells(session, citation, ['kb'], source_sha256=digest) == []
