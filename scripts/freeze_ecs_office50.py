"""Freeze 50 source-grounded Office questions before retrieval or generation."""
import hashlib,json,pathlib,re,sys,zipfile
from docx import Document
from openpyxl import load_workbook

base=pathlib.Path(sys.argv[1]);root=base/'office50-v1';sources=root/'sources';target=root/'dataset250'
if target.exists():raise RuntimeError('Frozen dataset cannot be overwritten')
target.mkdir(mode=0o700)
digest=lambda b:hashlib.sha256(b).hexdigest()
def clean(text):return re.sub(r'\s+',' ',text).strip()
rows=[];word_titles=[];word_values=[];word_ids=[]
for path in sorted(sources.glob('*.docx')):
 with zipfile.ZipFile(path) as z:
  if any(n.startswith('word/media/') for n in z.namelist()):continue
 doc=Document(path)
 if len(doc.tables)<2 or len(doc.tables[0].rows)<2 or len(doc.tables[1].rows)<2:continue
 h=[clean(c.text) for c in doc.tables[1].rows[0].cells]
 if h[:3]!=['Total Project Budget','Total Cost to Date for Project','Percentage of Work Completed to Date']:continue
 title=clean(doc.tables[0].rows[1].cells[0].text)
 if not title.lower().startswith('project title:'):continue
 title=title.split(':',1)[1].strip();source=digest(path.read_bytes());short=source[:12]
 def entry(question,answer,loc):
  return {'source_sha256':source,'source_path':str(path),'format':'docx','question':question,'answer':answer,
    'expected_sources':[{'sha256':source,**loc}],'kb_scope':['office50-v1'],'answer_points':[answer]}
 if title:word_titles.append(entry(f'In the transportation report {short}, what is the project title?',title,{'table':0,'row':1,'column':0}))
 vals=[clean(c.text) for c in doc.tables[1].rows[1].cells]
 for col,(header,value) in enumerate(zip(h,vals)):
  if value and len(value)<=55 and value.lower() not in ('na','n/a'):
   word_values.append(entry(f'In the transportation report {short}, what is the {header.lower()}?',value,{'table':1,'row':1,'column':col}))
 first=clean(doc.tables[0].rows[0].cells[0].text);m=re.search(r'TPF-\d\(\d+\)',first)
 if m:word_ids.append(entry(f'In the transportation report {short}, what is the project number?',m.group(),{'table':0,'row':0,'column':0}))
word=(word_titles+word_values+word_ids)[:25]
assert len(word)==25,('Not enough source-grounded Word questions',len(word))
specs=[]
for n in range(11,16):specs.append(('excel-1.xlsx','T01',n,2,'How many public libraries are recorded for {row} in table T01?'))
for n in range(9,14):specs.append(('excel-2.xlsx','T07',n,3,'What is the total public library operating revenue, in thousands, for {row} in table T07?'))
for c in range(2,7):specs.append(('excel-3.xlsx','Corn',5,c,'What percentage of all corn planted in {row} was insect-resistant Bt only in {col}?'))
for row,col,label in [(3,2,'average retail price per pound'),(4,2,'average retail price per pound'),(3,7,'average price per cup equivalent'),(4,7,'average price per cup equivalent'),(6,2,'average retail price per pint')]:
 specs.append(('excel-4.xlsx','Apples',row,col,'What is the '+label+' for the {row} form of apples?'))
for col,label in [(2,'average retail price per pound'),(4,'preparation yield factor'),(5,'size of a cup equivalent in pounds'),(7,'average price per cup equivalent')]:
 specs.append(('excel-5.xlsx','Avocado',3,col,'What is the '+label+' for fresh avocados?'))
specs.append(('excel-3.xlsx','Corn',6,2,'What percentage of all corn planted in {row} was insect-resistant Bt only in {col}?'))
assert len(specs)==25
excel=[]
for name,sheet,row,col,template in specs:
 path=sources/name;wb=load_workbook(path,read_only=True,data_only=True);ws=wb[sheet]
 value=ws.cell(row,col).value;row_label=clean(str(ws.cell(row,1).value or ''));col_label=clean(str(ws.cell(4,col).value or ''))
 assert isinstance(value,(int,float)) and row_label
 question=template.format(row=re.sub(r'\d+$','',row_label),col=col_label)
 source=digest(path.read_bytes());answer=str(value)
 excel.append({'source_sha256':source,'source_path':str(path),'format':'xlsx','question':question,'answer':answer,
  'numeric_answer':value,'numeric_relative_tolerance':0.0001,'answer_points':[answer],
  'expected_sources':[{'sha256':source,'sheet':sheet,'cell':ws.cell(row,col).coordinate}],'kb_scope':['office50-v1']})
 wb.close()
cases=word+excel
for i,row in enumerate(cases):row.update(qid=f'office50:{i+1:03d}',ordinal=204+i,dataset_version='office50-v1',expected_chunk_ids=[])
old=[]
for name in ('gold100.jsonl','qrel100.jsonl'):
 old.extend(json.loads(l) for l in (base/'dataset200'/name).read_text(encoding='utf-8').splitlines())
assert len(old)==200 and len({x['question'].strip().casefold() for x in old+cases})==250
payload=''.join(json.dumps(x,ensure_ascii=False,sort_keys=True)+'\n' for x in cases).encode()
(target/'office50.jsonl').write_bytes(payload)
oldhash=digest((base/'dataset200/freeze.json').read_bytes())
manifest={'status':'FROZEN_BEFORE_OFFICE_MODEL_INFERENCE','old200_freeze_sha256':oldhash,'office50_sha256':digest(payload),
 'total_unique_questions':250,'office_word':25,'office_excel':25,'office_answer_and_source_qrels':50,
 'gold_provenance':'deterministic extraction from original Office cells; generated questions, not official benchmark QA',
 'office_pdf_table_effect':'format controls: Word and Excel use same parser in both arms',
 'sources':[{'path':str(p),'sha256':digest(p.read_bytes())} for p in sorted({pathlib.Path(r['source_path']) for r in cases})]}
(target/'freeze250.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'status':manifest['status'],'word':25,'excel':25,'total':250,'source_count':len(manifest['sources']),'sha256':manifest['office50_sha256']}))
